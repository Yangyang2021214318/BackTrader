# -*- coding = utf-8 -*-
# @Time: 2026/9/28 15:41
# @Author: Yang Yang
# @File: inspect_columns
# @Software: PyCharm

from pathlib import Path
from openpyxl import load_workbook

from scripts.audit.inspect_data import RAW_DIR

RAW_DIR = Path(r"D:\A股日线回测\raw_data")
base_headers = None

for year in  (2015,2019,2025):
    path = sorted((RAW_DIR/str(year)).glob("*.xlsx"))[0]
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active

    rows = sheet.iter_rows(min_row=1, max_row=2, values_only=True)
    headers = next(rows)
    first_record = next(rows)

    print(f"\n{year}：{path.name}")
    print(f"工作表：{sheet.title}")
    print(f"字段数：{len(headers)}，数据行数：{sheet.max_row - 1}")

    print(f"日期示例：{first_record[0]!r}"
          f"类型：{type(first_record[0]).__name__}"
          )

    print(
        f"股票代码示例：{first_record[1]!r}"
        f"类型：{type(first_record[1]).__name__}"
    )

    if base_headers is None:
        base_headers = headers
        print("\n字段")
        for number, name in enumerate(headers, start=1):
            print(f"{number:2}.{name}")

    else:
        print(f"与 2015 年表头一致：{headers == base_headers}")

